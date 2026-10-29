"""汇编最终三文档；复用既有版式，完整收入设计分册，实测与未验证分开。"""
from pathlib import Path
import importlib.util
import json
import re
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables/最终提交'
spec = importlib.util.spec_from_file_location('course_doc', ROOT / 'scripts/build-course-documents.py')
base = importlib.util.module_from_spec(spec); spec.loader.exec_module(base)


def create(title, scope):
    doc = base.document(title, scope)
    doc.paragraphs[1].text = '2026年10月29日  课程交付版'
    doc.styles['Normal'].paragraph_format.space_after = Pt(4)
    return doc


def source(name):
    return (ROOT / 'docs' / name).read_text(encoding='utf-8-sig')


def section(doc, title, text):
    doc.add_page_break()
    doc.add_heading(title, 1)
    # 分册内的一级标题归入当前章节，保留原编号供源码追溯。
    text = re.sub(r'(?m)^## ', '### ', text)
    # 保留SQL/JSON/类型定义；只有Mermaid图仍通过随包原件查看。
    def fence(match):
        return match[0] if match[1].strip() == 'mermaid' else match[2]
    text = re.sub(r'```([^\n]*)\n(.*?)```', fence, text, flags=re.S)
    base.markdown(doc, text)


def save(doc, filename):
    # 表头与首行数据一起分页，避免页底只剩表头。
    for table in doc.tables:
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    doc.save(OUT / filename)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    req = create('阅集图书商城企业级需求文档', '面向课程评审、开发和测试，规定多商家实体新书商城的范围、角色权限、业务规则与验收目标。设计目标不等于已取得的实测结果；实际交付状态以测试报告为准。')
    base.markdown(req, source('01-总需求文档.md'))
    section(req, '交付范围及需求追踪', source('18-最终交付说明.md'))
    save(req, '企业级需求文档.docx')

    design = create('阅集图书商城设计文档', '完整汇编系统、数据库、接口、前端和部署运维设计。采用Java21、Spring Boot3.5.16模块化单体与独立Vue3前端，数据库迁移至V11。字段以03、网络契约以04及OpenAPI0.7.0为准。')
    design.add_paragraph('结构图的Mermaid原件及机器可读OpenAPI随源码包提供。文中历史阶段占位补充不作为当前接口实现状态；92项操作的实际观测见测试报告。冻结版本和镜像digest以infra锁文件及实测记录为准。')
    chapters = [('系统架构与一致性设计', '02-系统设计文档.md'), ('数据库与数据约束', '03-数据库设计.md'),
                ('接口契约与处理流程', '04-接口设计.md'), ('前端页面与交互', '05-前端页面与交互设计.md'),
                ('开发环境与部署', '06-开发环境与部署设计.md'), ('持续集成与运维', '08-CI-CD与运维设计.md')]
    for title, name in chapters:
        text = source(name)
        if name.startswith('04-'): text = text.split('## 6 阶段 A 技术占位补充')[0]
        text = '\n'.join(line for line in text.splitlines() if not (line.startswith(('阶段 ', '实施注记')) and '实施注记' in line))
        text = text.replace('；本次仅编写 docs', '').replace('未来代码、数据库迁移和运维文件均不得藏在设计文档目录中。', '代码、迁移和运维文件分别位于backend、frontend、infra及scripts。')
        section(design, title, text)
    save(design, '设计文档.docx')

    report = create('阅集图书商城测试报告', '依据实际报告及支付宝沙箱应答汇总。本地业务质量门槛、沙箱付款、有效通知和全额退款已有证据；未付款关单未确认，依赖漏洞未修复，其余外部演练按用户要求不再追加。本文不声明全部测试通过或已上线。')
    base.markdown(report, source('19-最终测试报告.md'))
    section(report, '测试设计及原验收口径', source('07-测试与验收设计.md'))
    save(report, '测试报告.docx')
    print(json.dumps({'outputDirectory': str(OUT), 'documents': [p.name for p in OUT.glob('*.docx')]}, ensure_ascii=False))


if __name__ == '__main__': main()
