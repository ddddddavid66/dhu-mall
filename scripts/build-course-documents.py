"""使用工作区提供的 python-docx 汇编课程三文档；不把目标或缺失项写成实测。"""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
FONT = 'Microsoft YaHei'


def plain(text):
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text.replace('**', '').replace('`', '').replace('<br>', '；')


def set_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string('111111')
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)


def document(title, scope):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    set_font(doc.styles['Normal'], 10.5)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.2
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    for style, size in [('Title', 23), ('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)]:
        set_font(doc.styles[style], size, True)
        doc.styles[style].paragraph_format.keep_with_next = True
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph('2026年10月26日  课程提交预览')
    doc.add_paragraph(scope)
    doc.add_paragraph('成员姓名、学号、班级及本人签名：待提交者补充。教师模板及截止时间尚未提供。')
    header = section.header.paragraphs[0]
    header.text = '阅集图书商城  ' + title
    set_font(doc.styles['Header'], 8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    footer.add_run(' 页')
    return doc


def add_table(doc, rows):
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=columns)
    table.style = 'Table Grid'
    table.autofit = False
    widths = [16.6 / columns] * columns
    if columns >= 3:
        widths = [2.2] + [14.4 / (columns-1)] * (columns-1)
    if rows[0][0].strip().startswith('实际命令'):
        widths = [5.4, 5.6, 5.6]
    for col, width in zip(table.columns, widths): col.width = Cm(width)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for column, cell in enumerate(cells):
            cell.width = Cm(widths[column])
            cell.text = plain(row[column]) if column < len(row) else ''
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5); run.bold = index == 0
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); table.rows[index]._tr.get_or_add_trPr().append(repeat)
        avoid_split = OxmlElement('w:cantSplit'); table.rows[index]._tr.get_or_add_trPr().append(avoid_split)
    doc.add_paragraph()


def markdown(doc, text):
    lines = text.splitlines(); index = 0; in_code = False
    while index < len(lines):
        line = lines[index].strip(); index += 1
        if line.startswith('```'):
            in_code = not in_code
            if in_code: doc.add_paragraph('结构关系图详见随包 Markdown 原件，可通过 Mermaid 查看。')
            continue
        if in_code or not line or line == '---' or line.startswith('<'): continue
        if line.startswith('|'):
            rows = [line.strip('|').split('|')]
            while index < len(lines) and lines[index].strip().startswith('|'):
                row = lines[index].strip().strip('|').split('|'); index += 1
                if all(re.fullmatch(r'\s*:?-+:?\s*', cell) for cell in row): continue
                rows.append([cell.strip() for cell in row])
            add_table(doc, rows); continue
        match = re.match(r'^(#{1,4})\s+(.+)', line)
        if match:
            if len(match[1]) == 1: continue
            heading = re.sub(r'[^\w\s\u4e00-\u9fff]', ' ', plain(match[2]))
            doc.add_heading(heading, min(len(match[1])-1, 3)); continue
        doc.add_paragraph(plain(line))


def main():
    OUT.mkdir(exist_ok=True)
    requirements = document('阅集图书商城需求文档', '本文汇编已接受的总需求基线。功能与非功能条款是验收目标；实际完成情况以测试报告为准。')
    markdown(requirements, (ROOT / 'docs/01-总需求文档.md').read_text(encoding='utf-8-sig'))
    requirements.add_heading('附录 需求追踪', 1)
    requirements.add_paragraph('需求与数据、API、UI、TC的逐项关系见随包 docs/11-实验要求追踪矩阵.md；接口92项与测试观测见 docs/stages/R-契约验收矩阵.md。')
    requirements.save(OUT / '01-阅集商城需求文档.docx')

    design = document('阅集图书商城设计文档', '本文汇编系统架构与一致性设计，并提供数据库、接口、页面、部署和CI的分册索引。Java21与Spring Boot3.5技术基线保持不变。')
    source = (ROOT / 'docs/02-系统设计文档.md').read_text(encoding='utf-8-sig')
    source = source.replace('；本次仅编写 docs', '').replace('未来代码、数据库迁移和运维文件均不得藏在设计文档目录中。', '代码、迁移和运维文件分别位于backend、frontend、infra及scripts。')
    markdown(design, source)
    design.add_heading('数据库与契约分册', 1)
    add_table(design, [
        ['领域', '关键对象及约束', '事实来源'],
        ['身份', '用户、角色权限、店铺成员、地址、轮换会话；登录范围由服务端推导', '03第3节；04 API01—10/70—77/84—88'],
        ['目录库存', 'Book独立于店铺Product/SKU；库存流水、预占不被人工调整覆盖', '03第4—5节；API11—25/63—68/90/92'],
        ['交易支付', '批次拆店订单、行/地址/金额快照；每订单唯一paymentNo，金额十进制字符串', '03第6—7节；API26—35'],
        ['履约售后', '物流状态与退款状态分离；退款成功或可再售验收分别恰好回补一次', '03第8节；API36—62'],
        ['结算支撑', '逐订单手续费快照、职责分离审核、模拟出款；审计、幂等、Outbox和消费去重', '03第9—10节；API78—83；V11'],
    ])
    design.add_paragraph('全部字段、外键和唯一约束以03及backend/src/main/resources/db/migration的V1—V11为准；92项网络模型及错误码以04和contracts/openapi.json 0.7.0为准。refund.reason_type按已批准ADR-17扩为VARCHAR(32)。')
    design.add_heading('前端与工程化设计索引', 1)
    add_table(design, [
        ['分册', '设计及实际边界'],
        ['05 前端', '独立Vue3工程；会员和工作台布局，UI01—33；统一请求层、合并刷新、幂等重试、切店取消旧请求；服务器金额和后端支付状态'],
        ['06 环境', '版本与digest冻结；dev/test/prod隔离；密钥受限文件注入；MySQL、Redis、RabbitMQ与持久文件卷'],
        ['07 质量', 'MySQL集成、契约、浏览器、安全并发与故障；固定80%服务分支门槛；TC01—32状态见实测报告'],
        ['08 运维', 'PR不访问部署秘密；SHA/digest晋级；兼容迁移、单机维护窗口、加密备份及隔离恢复；监控日志告警'],
        ['12 决策', 'ADR-19接受本地安全风险并取消兼容升级；不构成生产授权或扫描通过'],
    ])
    design.save(OUT / '02-阅集商城设计文档.docx')

    report_source = (ROOT / 'docs/stages/R-验证与交接.md').read_text(encoding='utf-8-sig')
    report_text = ('# 阅集图书商城本地验收报告\n\n'
                   '日期：2026-10-26。报告依据07测试设计和真实执行记录。Java21 / Boot3.5.16、V11、OpenAPI0.7.0；本机Windows加Docker Linux容器，MySQL隔离测试库，不是01规定的正式压测环境。\n\n'
                   '178项后端测试及80.1777%服务分支门槛通过，Vitest10项、桌面手机E2E5项通过，92项接口有成功与拒绝响应观测。第三方交易、标准性能、独立备份及完整24小时窗口仍有限制，不能称为全课程验收通过。17项镜像依赖发现按ADR-19接受本地风险，未进行兼容升级或生产发布。\n\n'
                   + report_source[report_source.index('## 实施与验证'):])
    (ROOT / 'docs/16-本地验收报告.md').write_text(report_text, encoding='utf-8')
    report = document('阅集图书商城测试报告', '本地业务质量门槛通过；真实沙箱买家、独立备份、正式压测、24小时窗口和部分远端验收仍有限制。本文按真实结果保留失败与阻塞。')
    markdown(report, report_text)
    report.save(OUT / '03-阅集商城测试报告.docx')
    print(json.dumps({'documents': [p.name for p in OUT.glob('*.docx')]}, ensure_ascii=False))


if __name__ == '__main__': main()
